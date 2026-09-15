"""Correct text in the supplied raster-based deck using native PPT overlays.

Requires python-pptx for artifact editing, not for running the app.
Usage: python src/tools/update_existing_deck.py --source /path/to/original.pptx
"""
import argparse
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

parser = argparse.ArgumentParser()
parser.add_argument('--source', type=Path, required=True)
parser.add_argument('--output', type=Path)
args = parser.parse_args()
root = Path(__file__).resolve().parents[2]
prs = Presentation(args.source)
N, W = '071A34', 'FFFFFF'
sx, sy = prs.slide_width / 1376, prs.slide_height / 768

def overlay(n, x, y, w, h, value, size=26, color=N, background=W, bold=False, center=False, link=None):
    slide = prs.slides[n-1]
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(x*sx), int(y*sy), int(w*sx), int(h*sy))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(background)
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.025)
    tf.margin_top = tf.margin_bottom = Inches(.025)
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
        p.space_after = Pt(0)
        p.line_spacing = 1.06
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size = 'Arial', Pt(size)
        r.font.bold, r.font.color.rgb = bold, RGBColor.from_string(color)
        if link:
            r.hyperlink.address = link
    return shape

# Preserve the supplied artwork and correct only content that disagrees with the app.
overlay(1,55,370,720,135,'Congestion forecasting and scenario\nsimulation for informed port planning.\nSynthetic-data prototype.',31)
overlay(1,55,632,850,103,'IBM Bobathon 2026  |  Team: The Watson Four\nLead: Deep Makwana\nMembers: Manav Kansagara, Smit Kansagara, Trupesh Hingrajiya',17,bold=True)
overlay(3,877,336,430,45,'Scenario congestion indicators',26,color='22A6B4',bold=True)
overlay(4,55,718,1240,30,'Illustrative workflow mockup — see the live demo for actual application output.',19,color='526477')
overlay(4,55,260,295,108,'• Combines synthetic\n  schedules with ML.',25)
overlay(4,55,372,295,94,'• Visualizes berth and\n  crane availability.',25)
overlay(4,55,460,295,160,'• Supports informed\n  planning decisions.',25)
overlay(5,60,60,1245,215,'The proposed planning journey\ntakes five steps',64,bold=True)
overlay(5,84,470,224,132,'Current input:\nseeded synthetic\nschedules and resources.',21)
overlay(5,326,470,223,132,'API validation and\nfeature preparation\nare implemented.',21)
overlay(5,570,470,223,132,'Trained models predict\nwaiting time and\ncongestion probability.',21)
overlay(5,813,470,225,142,'Roadmap: CP-SAT\nassignments and live\nBob explanations.\nCurrent UI is static.',21)
overlay(5,1055,470,242,142,'Demo approval changes\nlocal browser state.\nPersisted operational\nplans are future work.',21)
overlay(6,55,40,1250,139,'From working predictions\nto planned operational action',48,bold=True)
overlay(6,93,340,190,175,'Synthetic vessel\nschedules + berth\nand crane capacity',20,color=W,background='062242',center=True)
overlay(6,425,339,191,198,'scikit-learn\nGradient Boosting\nmodels: congestion\nprobability and\nvessel waiting time',20,background='08D2E9',center=True)
overlay(6,755,243,207,80,'Planning\n(roadmap)',28,color=W,background='2B679E',bold=True,center=True)
overlay(6,755,344,207,186,'Planned: OR-Tools\nCP-SAT berth and\ncrane assignments.\nCurrent interface:\nstatic demo.',22,color=W,background='2B679E',center=True)
overlay(6,1093,318,202,209,'Bob used during\ndevelopment.\nLive MCP explanations\nare planned.\nCopilot responses\nare predefined.',22,background='08D2E9',center=True)
overlay(6,80,621,1218,75,'Current: scenario calculations and trained prediction inference.\nNext: runtime scheduling and live IBM Bob MCP explanations.',28,center=True)
overlay(7,53,35,1280,81,'Design principle: models predict. Supervisors decide.',43,bold=True)
overlay(7,960,180,371,59,'Current build & roadmap',24,bold=True)
overlay(7,961,285,369,92,'Trained models generate\nwaiting-time and\ncongestion predictions.',23)
overlay(7,961,433,369,109,'IBM Bob used in development;\nlive MCP explanations\nare planned.',23)
overlay(7,961,590,369,112,'Supervisor approval is a\nlocal demo workflow.\nNo equipment is dispatched.',26)
overlay(8,51,39,1260,60,'Planning workflow preview — static demonstration',38,bold=True)
overlay(8,52,119,1260,64,'Illustration of the proposed planning and explanation workflow.\nNot runtime solver output or measured operational savings.',25)
overlay(8,713,225,314,36,'Copilot preview (not connected)',18,bold=True)
overlay(8,720,272,305,145,'Illustrative explanation only.\nLive IBM Bob MCP is not\nconnected, and a runtime\nsolver has not generated\nthese assignments.',22,background='CDF5F5')
overlay(8,757,445,267,53,'Example text only; no\noperational benefit measured.',16,background='E7EFF9')
overlay(8,1103,262,230,103,'Future: alternate routing\nrequires a runtime solver\nand cost comparison.',18)
overlay(8,1103,425,230,109,'Synthetic waiting-time\nestimates need real-port\nuncertainty validation.',18)
overlay(8,1103,600,230,102,'Demo approval requires\nconfirmation but changes\nlocal UI state only.',17)
overlay(9,61,50,1250,72,'Prototype architecture: working services and roadmap',39,bold=True)
overlay(9,532,248,372,43,'• React, Vite, TypeScript, Tailwind\n• Custom charts and static berth map',16,background='20CCE2')
overlay(9,535,382,353,102,'• Python, FastAPI, Pydantic\n• scikit-learn Gradient Boosting, NumPy\n• OR-Tools scheduling: planned\n• IBM Bob development; live MCP: planned',17,color=W,background='326BA4')
overlay(9,535,550,365,63,'• Optional PostgreSQL, SQLAlchemy, Alembic\n• Seeded synthetic schedules and resources\n• No confidential client data',16,color=W,background='10294A')
overlay(10,286,69,900,123,'Potential impact: earlier visibility\nand informed planning',44,bold=True,center=True)
overlay(10,361,257,207,77,'Baseline\nrules',29,color=W,background='2D689D',bold=True,center=True)
overlay(10,922,254,370,83,'72-hour planning\nvisibility',32,color=W,background='071A34',bold=True,center=True)
overlay(10,104,527,335,192,'Explore 72-hour synthetic\nscenarios and inspect\nupcoming resource pressure.\nReal-port validation is\nstill required.',26)
overlay(10,525,527,336,192,'Inspect berth and crane\ncapacity. Runtime CP-SAT\nis a next step. No measured\nturnaround savings.',24)
overlay(10,944,527,336,192,'Explicit confirmation is\ndemonstrated in the UI.\nNo automatic equipment\ndispatch or operational\nexecution.',26)
overlay(11,62,49,1254,194,'PortFlow AI: working\nforecasting prototype',56,bold=True)
overlay(11,137,297,525,40,'Architecture and reproducible setup',23)
overlay(11,137,385,525,40,'React + FastAPI dashboard and scenarios',21)
overlay(11,137,472,525,40,'Trained waiting-time and congestion models',21)
overlay(11,137,562,525,45,'IBM Bob development use; evidence pending',21)
overlay(11,770,297,524,40,'Local demo: localhost:5173 (not deployed)',22,link='http://localhost:5173')
overlay(11,770,385,524,40,'Video: recording and upload pending',23)
overlay(11,770,472,524,40,'Repo: github.com/Trupesh29/PortFlow-AI',22,link='https://github.com/Trupesh29/PortFlow-AI')
overlay(11,770,560,524,42,'Team: The Watson Four | Lead: Deep Makwana',20)
overlay(11,72,636,1230,31,'Pending: real video URL, Bob evidence, and repository naming/template confirmation.',18,center=True)

notes = {
1:'Friend: introduce PortFlow AI and the confirmed team. Synthetic-data forecasting prototype.',
2:'Friend: explain interacting arrivals, handling capacity, and vessel delays. No measured savings claimed.',
3:'Friend: contrast fragmented monitoring with scenario visibility and trained predictions.',
4:'Friend: identify this image as an illustrative workflow mockup, not the running application.',
5:'Friend: proposed journey. Uploads, runtime CP-SAT, live Bob MCP, and persisted approvals are not implemented.',
6:'Friend: working inference versus planned solver and MCP. Hand over for the actual live demo. See demo/recording-script.md.',
7:'Friend after demo: describe team-reported Bob development use and supervisor control. Exact tasks/evidence still needed.',
8:'Optional static planning preview. Diagram values are illustrative, not computed savings. Skip in a short recording.',
9:'Architecture: React/FastAPI, trained Gradient Boosting, optional database, planned solver/MCP.',
10:'Potential impact only; real-port validation and runtime solver are next steps.',
11:'Closing: show actual local/repository links. Video upload and submission metadata/evidence remain pending.'}
for n, note in notes.items():
    prs.slides[n-1].notes_slide.notes_text_frame.text = note
out = args.output or root / 'presentation/slides.pptx'
prs.save(out)
print(f'Updated existing deck: {out}; retained {len(prs.slides)} original slide images')
